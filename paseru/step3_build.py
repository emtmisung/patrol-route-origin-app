"""3단계 노선 생성 · 결과.

원래 app.py 한 파일에 있던 화면 코드를 그대로 옮기고 render(ctx) 함수로 감쌌다."""
import io
import math
import re
import zipfile
from datetime import datetime
from urllib.parse import quote

import folium
import pandas as pd
import streamlit as st
from streamlit_folium import st_folium

from paseru.exports import (
    KAKAO_MAX_ROUTE_POINTS,
    build_center_route_print_html,
    build_printable_qr_html,
    build_qr_zip,
    build_route_links_excel,
    kakao_route_links,
    kakao_url,
    make_qr_png,
)
from paseru.geo import has_keys, haversine_km, road_route
from paseru.maps import stop_label
from paseru.routing import (
    allocate_hydrants_by_distribution,
    build_routes,
    real_leg,
    separate_long_distance,
    separate_long_time,
)
from paseru.settings import (
    API_CALL_LIMIT,
    AVG_SPEED_KMH,
    BROWSER_DRAFT_DAYS,
    ROAD_FACTOR,
    app_public_url,
)
from paseru.storage import create_mobile_transfer, minimum_transfer_targets, route_execution_content
from paseru.ui import inspection_capacity_warning, route_generation_spinner, safety_warning
from paseru.upload import find_address_column_index, find_name_column_index


def render(ctx):
    # 앞 단계에서 만든 값(원래 한 파일의 전역 변수였던 것)
    basis = ctx.get("basis")
    basis_label = ctx.get("basis_label")
    candidate_k = ctx.get("candidate_k")
    commander_count = ctx.get("commander_count")
    commander_oneway_limit = ctx.get("commander_oneway_limit")
    commander_route_count = ctx.get("commander_route_count")
    commander_stop_min = ctx.get("commander_stop_min")
    coord_api_calls = ctx.get("coord_api_calls")
    df = ctx.get("df")
    end_dt = ctx.get("end_dt")
    exc = ctx.get("exc")
    guard_repeat_label = ctx.get("guard_repeat_label")
    guard_rounds = ctx.get("guard_rounds")
    hydrant_distribution_basis = ctx.get("hydrant_distribution_basis")
    hydrant_inspection_min = ctx.get("hydrant_inspection_min")
    hydrant_max_min = ctx.get("hydrant_max_min")
    hydrant_member_count = ctx.get("hydrant_member_count")
    hydrant_target_min = ctx.get("hydrant_target_min")
    hydrant_vehicle_count = ctx.get("hydrant_vehicle_count")
    hydrant_workdays = ctx.get("hydrant_workdays")
    inspect_capacity = ctx.get("inspect_capacity")
    inspect_dates = ctx.get("inspect_dates")
    inspect_minutes = ctx.get("inspect_minutes")
    inspect_targets_per_day = ctx.get("inspect_targets_per_day")
    inspect_teams = ctx.get("inspect_teams")
    lat = ctx.get("lat")
    lng = ctx.get("lng")
    long_threshold = ctx.get("long_threshold")
    max_calls = ctx.get("max_calls")
    max_per_route = ctx.get("max_per_route")
    max_routes_cap = ctx.get("max_routes_cap")
    mode = ctx.get("mode")
    now_timestamp = ctx.get("now_timestamp")
    origin = ctx.get("origin")
    patrol_title = ctx.get("patrol_title")
    period_days = ctx.get("period_days")
    period_end = ctx.get("period_end")
    period_start = ctx.get("period_start")
    purpose = ctx.get("purpose")
    purpose_label = ctx.get("purpose_label")
    qr_box = ctx.get("qr_box")
    qr_png = ctx.get("qr_png")
    remaining_seconds = ctx.get("remaining_seconds")
    route_prefix = ctx.get("route_prefix")
    season_actor = ctx.get("season_actor")
    season_delegate = ctx.get("season_delegate")
    season_limit_basis = ctx.get("season_limit_basis")
    season_oneway_limit = ctx.get("season_oneway_limit")
    season_stop_min = ctx.get("season_stop_min")
    season_target_min = ctx.get("season_target_min")
    season_vehicle = ctx.get("season_vehicle")
    seg_max_km = ctx.get("seg_max_km")
    seg_max_min = ctx.get("seg_max_min")
    stale_key = ctx.get("stale_key")
    start_dt = ctx.get("start_dt")
    station_address = ctx.get("station_address")
    station_lat = ctx.get("station_lat")
    station_lng = ctx.get("station_lng")
    station_name = ctx.get("station_name")
    station_query = ctx.get("station_query")
    station_result = ctx.get("station_result")
    target_min = ctx.get("target_min")
    target_min_high = ctx.get("target_min_high")
    total_calls = ctx.get("total_calls")
    transfer_expires_at = ctx.get("transfer_expires_at")
    transfer_token = ctx.get("transfer_token")
    transfer_url = ctx.get("transfer_url")
    vehicle = ctx.get("vehicle")

    # 출발지(소방서·센터) 자신이 순찰 대상 목록에 섞여 있으면 제외한다.
    # (업로드 파일 첫 줄에 소방서를 넣어두는 경우가 많아, 그대로 두면 소방서가 경유지로 잡힌다)
    excluded_station_rows = 0
    if df is not None and len(df):
        def _norm(v):
            return re.sub(r"\s+", "", str(v)) if v is not None else ""

        st_name_n, st_addr_n = _norm(station_name), _norm(station_address)
        mask_keep = []
        for _, r in df.iterrows():
            vals = [_norm(v) for v in r.values]
            is_station = any(v and (v == st_name_n or v == st_addr_n) for v in vals)
            mask_keep.append(not is_station)
        excluded_station_rows = len(df) - sum(mask_keep)
        if excluded_station_rows:
            df = df[pd.Series(mask_keep, index=df.index)].reset_index(drop=True)

    # ----------------------------------------------------------------------------
    # 좌표 확인·수정 표 + 노선 생성 버튼
    # (설정 요약·업로드 데이터 확인 등은 걷어내고, 좌표 수정과 실행 버튼만 남긴다)
    # ----------------------------------------------------------------------------
    if df is not None and len(df):
        cols = list(df.columns)
        name_col_guess_idx = find_name_column_index(cols)
        addr_col_guess_idx = find_address_column_index(cols, name_col_guess_idx)

        name_col = cols[name_col_guess_idx]
        addr_col = cols[addr_col_guess_idx]

        lat_col_guess = next((c for c in cols if "위도" in str(c) or str(c).lower() == "lat"), None)
        lng_col_guess = next((c for c in cols if "경도" in str(c) or str(c).lower() in ("lng", "lon")), None)
        has_coords = bool(lat_col_guess and lng_col_guess)
        coord_mode = "file" if has_coords else "geocode"

        n_targets = len(df)

        saved_coords = st.session_state.get("coords_df")
        coords_complete = (saved_coords is not None and len(saved_coords) == n_targets
                           and saved_coords["위도"].notna().all()
                           and saved_coords["경도"].notna().all())
        if st.session_state.get("coord_future") is not None:
            st.info("🔍 좌표 검색 중입니다. 1단계 기본정보에서 완료 상태를 확인하세요.")
        elif saved_coords is None:
            st.warning("⚠️ 1단계 기본정보에서 좌표를 먼저 검색해주세요.")

        coords_df = st.session_state.get("coords_df")

        if coords_df is not None:
            with st.container(border=True):
                st.markdown("### 🔎 좌표 확인 및 수정")
                ok_n = int(coords_df["위도"].notna().sum())
                fail_n = int(coords_df["위도"].isna().sum())
                k1, k2, k3 = st.columns(3)
                k1.metric("전체", f"{len(coords_df)}")
                k2.metric("좌표 확보", f"{ok_n}")
                k3.metric("좌표 없음", f"{fail_n}")

                if fail_n:
                    st.error(
                        f"❌ {fail_n}건은 좌표를 찾지 못했습니다. 1단계의 **좌표 실패건 처리**에서 "
                        "주소 수정 후 재검색·좌표 직접입력·대상 제외 중 하나를 선택하거나, "
                        "아래 표의 위도·경도 칸에 직접 입력하세요."
                    )
                else:
                    st.success("✅ 모든 대상의 좌표가 확보되었습니다. 아래에서 노선을 생성하세요.")

                st.caption("위도·경도 칸은 직접 고칠 수 있습니다. 수정하면 그 값이 노선 생성에 그대로 쓰입니다.")
                edited = st.data_editor(
                    coords_df, use_container_width=True, hide_index=True, num_rows="fixed",
                    key="coords_editor",
                    column_config={
                        "대상명": st.column_config.TextColumn(disabled=True, width="medium"),
                        "주소": st.column_config.TextColumn(disabled=True, width="large"),
                        "위도": st.column_config.NumberColumn(format="%.6f", help="예: 35.919000"),
                        "경도": st.column_config.NumberColumn(format="%.6f", help="예: 128.283000"),
                        "상태": st.column_config.TextColumn(disabled=True, width="small"),
                        "비고": st.column_config.TextColumn(disabled=True, width="large"),
                    },
                )
                # 사용자가 별도 저장 버튼을 누르지 않아도 수정된 좌표를 현재 작업에 즉시 보존한다.
                st.session_state["coords_df"] = edited.copy()

                st.markdown(
                    '''<div style="margin-top:0.75rem;padding:0.9rem 1rem;
                        border:1px solid #48a774;border-left:6px solid #238553;border-radius:10px;
                        background:#e5f6ed;color:#155f3b;font-weight:750;line-height:1.55;">
                        ✅ 확정된 좌표는 현재 작업 동안 백그라운드에 자동 저장되어 노선 생성에 바로 반영됩니다.<br>
                        <span style="font-weight:550;color:#28704b;">별도의 저장 버튼을 누르지 않아도 됩니다.</span>
                    </div>''',
                    unsafe_allow_html=True,
                )

            ready = edited["위도"].notna() & edited["경도"].notna()
            n_ready = int(ready.sum())
            if candidate_k:
                est_calls = n_ready * candidate_k + n_ready
            else:
                est_calls = n_ready * (n_ready + 1) // 2 + n_ready
            if n_ready < len(edited):
                safety_warning(
                    f"좌표가 없는 {len(edited) - n_ready}건은 노선에서 제외됩니다. "
                    "1단계의 좌표 실패건 처리에서 재검색·좌표 직접입력·제외 중 하나를 선택하세요.",
                    title="좌표 없는 대상 안내",
                )

            if purpose == "inspect":
                inspect_capacity = (
                    len(inspect_dates) * int(inspect_teams) * int(inspect_targets_per_day)
                )
                inspect_omitted_count = max(0, n_ready - inspect_capacity)
                capacity_formula = (
                    f"검사 가능일 {len(inspect_dates)}일 × {int(inspect_teams)}팀 × "
                    f"팀당 하루 {int(inspect_targets_per_day)}개소 = 최대 {inspect_capacity}개소"
                )
                if inspect_omitted_count:
                    inspection_capacity_warning(
                        capacity_formula,
                        n_ready,
                        inspect_omitted_count,
                    )
                elif inspect_capacity:
                    st.success(f"✅ {capacity_formula} · 전체 {n_ready}개소를 기간 안에 검사할 수 있습니다.")
                else:
                    st.error("검사 가능한 날짜가 없습니다. 검사기간 또는 검사 가능 요일을 조정하세요.")

            run_disabled = (
                not has_keys() or n_ready == 0 or
                (purpose == "inspect" and inspect_capacity == 0)
            )
            run_col1, run_col2 = st.columns(2)
            with run_col1:
                run = st.button(
                    "🚒 노선 생성 시작", type="primary",
                    disabled=run_disabled,
                    use_container_width=True,
                )
            with run_col2:
                rerun_same_condition = st.button(
                    "🔁 같은 조건으로 노선 재탐색",
                    type="secondary",
                    disabled=run_disabled,
                    help="주소·일정·팀 수는 그대로 두고 다른 후보 순서로 노선을 다시 탐색합니다.",
                    use_container_width=True,
                )
            if run:
                st.session_state["route_search_variant"] = 0
            if rerun_same_condition:
                st.session_state["route_search_variant"] = (
                    int(st.session_state.get("route_search_variant", 0)) + 1
                )
                for stale_key in ("route_results", "far_points", "meta"):
                    st.session_state.pop(stale_key, None)
                run = True
        else:
            run = False
            edited = None
            st.info("먼저 위의 좌표 확인 및 수정에서 좌표를 확정해 주세요.")
    else:
        run = False
        edited = None
        st.info("먼저 1단계 기본정보에서 대상 목록을 업로드해 주세요.")

    if run:
        route_variant = int(st.session_state.get("route_search_variant", 0))
        route_status = st.empty()
        route_generation_spinner(
            route_status,
            "노선 생성 중입니다",
            "돋보기가 움직이는 동안 실제 도로 기준 노선과 카카오맵 연결 정보를 만들고 있습니다.",
        )
        # ---- 중단 장치 ----------------------------------------------------
        # ① 수동 중단: 아래 '중단' 버튼을 누르면 Streamlit이 새로 실행되면서
        #    지금 돌고 있는 계산이 즉시 멈춘다.
        # ② 자동 중단: 호출 수가 한도(max_calls)를 넘으면 그때까지 만든 노선만 남기고 멈춘다.
        stop_box = st.container()
        with stop_box:
            st.button("⛔ 계산 중단", key="stop_btn", type="secondary",
                      help="지금 진행 중인 노선 생성을 즉시 멈춥니다.")

        call_counter = {"n": 0}
        limit_hit = {"v": False}

        def over_limit():
            if call_counter["n"] >= max_calls:
                limit_hit["v"] = True
                return True
            return False

        # 1) 기본정보에서 검색·확정한 출발부서 좌표 사용
        s_lat, s_lng = station_lat, station_lng
        if s_lat is None or s_lng is None:
            st.error("1단계 기본정보에서 출발부서를 검색하거나 현 위치를 조회해주세요.")
            st.stop()
        station = {"name": station_name, "lat": s_lat, "lng": s_lng}

        # 2) 1단계에서 확정한 좌표를 그대로 사용 (여기서는 지오코딩을 하지 않는다)
        points = []
        for _, r in edited.iterrows():
            try:
                lat, lng = float(r["위도"]), float(r["경도"])
            except (TypeError, ValueError):
                continue
            if math.isnan(lat) or math.isnan(lng):
                continue
            points.append({"name": str(r["대상명"]), "address": str(r["주소"]),
                           "lat": lat, "lng": lng})

        if not points:
            st.error("좌표가 있는 대상이 없습니다. 좌표 확인 및 수정에서 좌표를 확정해 주세요.")
            st.stop()

        # 3) 용도별 원거리 판정
        if purpose == "hydrant":
            normal_points = allocate_hydrants_by_distribution(
                points, station, int(hydrant_vehicle_count), hydrant_distribution_basis
            )
            far_points = []
        elif purpose == "inspect":
            # 예방검사는 거리에 관계없이 업로드한 모든 대상을 반드시 포함한다.
            normal_points = points
            far_points = []
        elif purpose == "other":
            long_progress = st.progress(0.0, text="출발지 기준 편도시간 확인 중...")

            def bump(total_hint=len(points)):
                call_counter["n"] += 1
                long_progress.progress(
                    min(call_counter["n"] / max(total_hint, 1), 1.0),
                    text=f"실제 도로시간 API 호출 중... ({call_counter['n']:,}/{max_calls:,}회)",
                )

            normal_points, far_points = separate_long_time(
                points, station, float(commander_oneway_limit),
                "편도 허용시간 초과 - 별도 지휘구역 검토",
                on_call=bump, should_stop=over_limit,
            )
            long_progress.empty()
        else:
            long_progress = st.progress(0.0, text="소방서 기준 실도로거리 확인 중...")

            def bump(total_hint=len(points)):
                call_counter["n"] += 1
                long_progress.progress(min(call_counter["n"] / max(total_hint, 1), 1.0),
                                       text=f"실제 도로거리 API 호출 중... "
                                            f"({call_counter['n']:,}/{max_calls:,}회)")

            if purpose == "season":
                if season_limit_basis == "편도시간(분)":
                    normal_points, far_points = separate_long_time(
                        points, station, float(season_oneway_limit), season_delegate,
                        on_call=bump, should_stop=over_limit,
                    )
                else:
                    normal_points, far_points = separate_long_distance(
                        points, station, float(season_oneway_limit), on_call=bump,
                        save_calls=False, should_stop=over_limit,
                    )
                    far_points = [{**point, "권장수행": season_delegate} for point in far_points]
            else:
                normal_points, far_points = separate_long_distance(
                    points, station, long_threshold, on_call=bump, save_calls=bool(candidate_k),
                    should_stop=over_limit,
                )
            long_progress.empty()

        # 4) 노선 편성
        build_progress = st.empty()

        def bump_build():
            call_counter["n"] += 1
            if purpose == "hydrant":
                build_progress.text("지리조사 노선을 자동 편성하고 있습니다...")
            else:
                build_progress.text(f"실도로 기준 노선 편성 중... "
                                    f"(API 호출 {call_counter['n']:,}/{max_calls:,}회)")

        if purpose == "hydrant":
            routes = []
            unassigned = []
            for vehicle_no in range(1, int(hydrant_vehicle_count) + 1):
                vehicle_points = [p for p in normal_points if p.get("vehicle_no") == vehicle_no]
                if not vehicle_points:
                    continue
                vehicle_routes, vehicle_unassigned = build_routes(
                    vehicle_points, station, "target_time", 100,
                    None, None, int(hydrant_target_min),
                    None, basis="time", on_call=bump_build,
                    candidate_k=candidate_k, should_stop=over_limit,
                    service_min_per_stop=int(hydrant_inspection_min),
                    route_variant=route_variant,
                )
                routes.extend(vehicle_routes)
                unassigned.extend(vehicle_unassigned)
        elif purpose == "other":
            commander_per_route = max(1, math.ceil(len(normal_points) / commander_route_count))
            routes, unassigned = build_routes(
                normal_points, station, "fixed", commander_per_route,
                None, None, None, commander_route_count,
                basis="distance", on_call=bump_build,
                candidate_k=candidate_k, should_stop=over_limit,
                service_min_per_stop=int(commander_stop_min),
                route_variant=route_variant,
            )
        else:
            routes, unassigned = build_routes(
                normal_points, station, mode, max_per_route,
                seg_max_km, seg_max_min, target_min_high,
                (max_routes_cap if purpose == "inspect" else max_routes_cap or None),
                basis=basis, on_call=bump_build,
                candidate_k=candidate_k, should_stop=over_limit,
                service_min_per_stop=(int(season_stop_min) if purpose == "season" else
                                      int(inspect_minutes) if purpose == "inspect" else 0),
                strict_route_cap=(purpose == "inspect"),
                route_variant=route_variant,
            )
        build_progress.empty()

        if limit_hit["v"]:
            if purpose == "hydrant":
                st.warning(
                    f"⛔ 계산 범위 한도에 도달해 편성을 중단했습니다. "
                    f"그때까지 편성된 {len(routes)}개 노선은 아래에 표시됩니다."
                )
            else:
                st.warning(
                    f"⛔ API 호출 한도({max_calls:,}회)에 도달해 노선 편성을 중단했습니다. "
                    f"그때까지 편성된 {len(routes)}개 노선은 아래에 그대로 표시됩니다. "
                    "한도를 늘리거나 'API 호출 절약'을 켜고 다시 실행해 보세요."
                )

        inspect_omitted_points = unassigned[:] if purpose == "inspect" else []
        if unassigned and purpose != "inspect":
            for p in unassigned:
                if over_limit():
                    km = haversine_km(
                        (station["lat"], station["lng"]), (p["lat"], p["lng"])
                    ) * ROAD_FACTOR
                else:
                    km, _ = real_leg(station, p, on_call=bump_build)
                far_points.append({**p, "도로거리_km": round(km, 1)})

        # 용도별 부가 정보
        team_info = ""
        if purpose == "hydrant":
            base_count, extra_count = divmod(len(points), max(int(hydrant_member_count), 1))
            count_text = (f"{base_count}~{base_count + 1}개" if extra_count else f"{base_count}개")
            team_info = (f" · {hydrant_member_count}명 개인별 {count_text}"
                         f" · 차량 {hydrant_vehicle_count}대 · 월 {hydrant_workdays}근무일"
                         f" · {hydrant_distribution_basis}")
        elif purpose == "season":
            limit_unit = "분" if season_limit_basis == "편도시간(분)" else "km"
            team_info = (f" · 관할 전체 대상 균등 순환 · {season_actor}/{season_vehicle}"
                         f" · 1회 최대 {int(season_target_min)}분"
                         f" · 편도 {season_oneway_limit:g}{limit_unit} 제한")
        elif purpose == "other":
            team_info = (
                f" · 지휘관 {int(commander_count)}명 · 자동 생성구역 {len(routes)}개"
                f" · 편도 {int(commander_oneway_limit)}분 이내"
                f" · 현장당 {int(commander_stop_min)}분"
            )
        elif purpose == "inspect":
            available_team_days = len(inspect_dates) * int(inspect_teams)
            assigned_targets = sum(len(route) for route in routes)
            team_info = (f" · 검사 가능일 {len(inspect_dates)}일 · {inspect_teams}팀"
                         f" · 팀당 하루 최대 {int(inspect_targets_per_day)}개소")
            if inspect_omitted_points:
                inspection_capacity_warning(
                    (f"검사 가능일 {len(inspect_dates)}일 × {int(inspect_teams)}팀 × "
                     f"팀당 하루 {int(inspect_targets_per_day)}개소 = 최대 {inspect_capacity}개소"),
                    len(points),
                    len(inspect_omitted_points),
                )
            else:
                st.success(
                    f"✅ 검사 가능일 {len(inspect_dates)}일, {int(inspect_teams)}팀, "
                    f"팀당 하루 최대 {int(inspect_targets_per_day)}개소 조건으로 "
                    f"전체 {assigned_targets}개소를 기간 안에 편성했습니다."
                )
        elif purpose == "guard" and guard_repeat_label == "매일 같은 코스 반복" and guard_rounds:
            total_runs = int(guard_rounds.replace("회", "")) * period_days
            team_info = f" · 매일 같은 코스로 하루 {guard_rounds} 반복({period_days}일간 총 {total_runs}회)"
        elif purpose == "guard":
            team_info = f" · 매일 다른 코스로 순환({period_days}일간 {len(routes)}개 노선 배정)"

        far_word = "편도 기준 초과" if purpose in ("season", "other") else "장거리 별도"
        far_summary = "" if purpose == "inspect" else f" ({far_word} {len(far_points)}개소)"
        st.success(f"[{purpose_label}] 총 {len(routes)}개 노선, {sum(len(r) for r in routes)}개소 배정 완료 "
                   f"{far_summary}{team_info}")
        # 5) 확정 노선의 구간별 실도로거리·경로좌표
        route_results = []
        total_calls = sum(len(r) + 1 for r in routes)
        call_progress = st.progress(0.0, text="노선별 실도로 경로 확정 중...")
        done = 0
        for ri, route in enumerate(routes):
            legs = []
            cur = station
            acc_km = 0.0
            acc_min = 0.0
            all_path = []
            for p in route:
                km, mins, path = road_route(cur["lat"], cur["lng"], p["lat"], p["lng"])
                if km is None:
                    km = haversine_km((cur["lat"], cur["lng"]), (p["lat"], p["lng"])) * ROAD_FACTOR
                    mins = km / AVG_SPEED_KMH * 60
                    path = [(cur["lat"], cur["lng"]), (p["lat"], p["lng"])]
                service_min = (int(hydrant_inspection_min) if purpose == "hydrant" else
                               int(season_stop_min) if purpose == "season" else
                               int(commander_stop_min) if purpose == "other" else 0)
                if purpose == "inspect":
                    service_min = int(inspect_minutes)
                legs.append({"from": cur["name"], "to": p["name"], "to_address": p.get("address", ""),
                             "km": km, "min": mins, "inspection_min": service_min,
                             "assigned_to": p.get("assigned_to", ""),
                             "vehicle_no": p.get("vehicle_no"),
                             "lat": p["lat"], "lng": p["lng"]})
                acc_km += km
                acc_min += mins + service_min
                all_path += path
                cur = p
                done += 1
                call_progress.progress(min(done / max(total_calls, 1), 1.0), text="노선별 실도로 경로 확정 중...")
            back_km, back_min, back_path = road_route(cur["lat"], cur["lng"], station["lat"], station["lng"])
            if back_km is None:
                back_km = haversine_km((cur["lat"], cur["lng"]), (station["lat"], station["lng"])) * ROAD_FACTOR
                back_min = back_km / AVG_SPEED_KMH * 60
                back_path = [(cur["lat"], cur["lng"]), (station["lat"], station["lng"])]
            acc_km += back_km
            acc_min += back_min
            all_path += back_path
            done += 1
            call_progress.progress(min(done / max(total_calls, 1), 1.0), text="노선별 실도로 경로 확정 중...")

            result = {
                "route_no": ri + 1, "stops": route, "legs": legs,
                "vehicle_no": route[0].get("vehicle_no") if route else None,
                "assigned_members": sorted({p.get("assigned_to", "") for p in route if p.get("assigned_to")}),
                "back_km": back_km, "back_min": back_min,
                "total_km": acc_km, "total_min": acc_min, "path": all_path,
            }
            if purpose == "inspect" and inspect_dates:
                date_index = ri // int(inspect_teams)
                if date_index < len(inspect_dates):
                    result["inspection_date"] = inspect_dates[date_index]
                    result["inspection_team"] = ri % int(inspect_teams) + 1
            route_results.append(result)
        call_progress.empty()

        if purpose == "hydrant":
            route_counts = {
                vehicle_no: sum(1 for result in route_results if result.get("vehicle_no") == vehicle_no)
                for vehicle_no in range(1, int(hydrant_vehicle_count) + 1)
            }
            over_vehicles = {v: count for v, count in route_counts.items() if count > int(hydrant_workdays)}
            if over_vehicles:
                detail = ", ".join(f"{v}호차 {count}개 노선" for v, count in over_vehicles.items())
                suggested = min(
                    int(hydrant_max_min),
                    math.ceil(int(hydrant_target_min) * max(over_vehicles.values()) / int(hydrant_workdays) / 10) * 10,
                )
                st.warning(
                    f"월 {int(hydrant_workdays)}번의 당번일을 초과하는 차량이 있습니다: {detail}. "
                    f"우선 목표시간을 약 {suggested}분으로 늘려 다시 편성해 보세요. "
                    f"{int(hydrant_max_min)}분에서도 10개 이내가 되지 않으면 인원 또는 차량 편성을 조정해야 합니다."
                )
            else:
                detail = " · ".join(f"{v}호차 {count}개 노선" for v, count in route_counts.items())
                st.success(f"월 {int(hydrant_workdays)}번의 당번일 안에 전수조사가 가능합니다. {detail}")
        elif purpose == "season":
            required_routes = len(route_results)
            if required_routes > period_days:
                st.warning(
                    f"전 대상을 누락 없이 순찰하려면 1회 최대 {int(season_target_min)}분 기준으로 "
                    f"최소 {required_routes}일이 필요하지만 "
                    f"설정한 기간은 {period_days}일입니다. 기간을 {required_routes}일 이상으로 늘리거나 "
                    "1회 최대시간을 늘려 다시 편성해주세요."
                )
            elif required_routes:
                base_runs, extra_runs = divmod(period_days, required_routes)
                for route_index, result in enumerate(route_results):
                    result["period_runs"] = base_runs + (1 if route_index < extra_runs else 0)
                min_runs = base_runs
                max_runs = base_runs + (1 if extra_runs else 0)
                run_text = f"{min_runs}회" if min_runs == max_runs else f"{min_runs}~{max_runs}회"
                st.success(
                    f"일반 순찰 대상 {sum(len(result['stops']) for result in route_results)}개소를 "
                    f"{required_routes}개 코스로 나누었습니다. {period_days}일 동안 코스를 차례로 순환하면 "
                    f"각 코스와 소속 대상은 예상 {run_text} 방문합니다."
                )
                if far_points:
                    st.info(
                        f"편도 기준을 넘는 {len(far_points)}개소는 출동 공백을 줄이기 위해 "
                        f"{season_delegate} 대상으로 별도 안내합니다."
                    )

        st.session_state["station"] = station
        st.session_state["route_results"] = route_results
        st.session_state["far_points"] = far_points
        route_status.empty()

        st.session_state["meta"] = {
            "title": patrol_title, "purpose": purpose_label, "vehicle": vehicle,
            "period": (f"{period_start:%Y-%m-%d} ~ {period_end:%Y-%m-%d} "
                       f"(검사 가능일 {len(inspect_dates)}일)" if purpose == "inspect" else
                       "" if purpose == "hydrant" else
                       f"{period_start:%Y-%m-%d}" if purpose == "other" else
                       f"{start_dt:%Y-%m-%d %H:%M} ~ {end_dt:%Y-%m-%d %H:%M} ({period_days}일간)"),
            "period_label": ("검사기간" if purpose == "inspect" else
                             "" if purpose == "hydrant" else
                             "방문일" if purpose == "other" else "순찰기간"),
            "basis": basis_label, "route_prefix": route_prefix, "team_info": team_info.strip(" ·"),
            "target_min": target_min,
            "hydrant_distribution_basis": hydrant_distribution_basis,
            "route_search_variant": route_variant,
            "api_calls_used": coord_api_calls + call_counter["n"] + total_calls,
            "api_call_limit": API_CALL_LIMIT,
            "validated_target_count": 217,
        }

    # ----------------------------------------------------------------------------
    # 결과 표시 (구 ⑤ 노선 결과 탭 내용 — 이제 이 탭 안에서 이어서 보여준다)
    # ----------------------------------------------------------------------------
    if st.session_state.get("stop_btn"):
        st.info("⛔ 계산을 중단했습니다. (중단 시점까지의 계산 결과는 저장되지 않습니다. "
                "직전에 완료된 결과가 있으면 아래에 그대로 남아 있습니다.)")

    if "route_results" in st.session_state:
        station = st.session_state["station"]
        route_results = st.session_state["route_results"]
        far_points = st.session_state["far_points"]
        meta = st.session_state.get("meta", {})

        st.write("")
        st.divider()
        st.header("📍 노선 생성 결과")
        if meta:
            period_summary = (f" · {meta.get('period_label')} {meta.get('period')}"
                              if meta.get("period_label") and meta.get("period") else "")
            st.caption(f"**{meta.get('title','')}** · {meta.get('purpose','')} · 기준: {meta.get('basis','')}"
                       + period_summary
                       + (f" · 차량: {meta['vehicle']}" if meta.get("vehicle") else "")
                       + (f" · {meta['team_info']}" if meta.get("team_info") else ""))
        if meta.get("purpose") == "① 지휘관 현장방문":
            safety_warning(
                "이 노선은 평시 순찰계획과 사전 검토용입니다. "
                "실제 재난현장에서는 실시간 상황이 반영되지 않으므로, "
                "본 계산만으로 지휘·대응을 결정하지 말고 현장지휘관의 판단과 공식 지휘체계를 우선하십시오."
            )

        metric_columns = st.columns(4 if meta.get("purpose") == "④ 예방검사" else 5)
        m1, m2, m3, m4 = metric_columns[:4]
        m1.metric("생성 노선 수", f"{len(route_results)}")
        m2.metric("전체 방문지", f"{sum(len(r['stops']) for r in route_results)}")
        m3.metric("총 이동거리(km)", f"{sum(r['total_km'] for r in route_results):.1f}")
        m4.metric("노선 평균시간", f"{sum(r['total_min'] for r in route_results) / max(len(route_results), 1):.0f}분")
        if meta.get("purpose") != "④ 예방검사":
            m5 = metric_columns[4]
            m5.metric("편도 기준 초과" if meta.get("purpose") in ("① 지휘관 현장방문", "③ 계절순찰") else "원거리 분리 대상",
                      f"{len(far_points)}")
        if meta.get("purpose") == "③ 계절순찰":
            visit_runs = [r.get("period_runs", 0) for r in route_results if r.get("period_runs")]
            if visit_runs:
                repeat_text = (f"{min(visit_runs)}회" if min(visit_runs) == max(visit_runs)
                               else f"{min(visit_runs)}~{max(visit_runs)}회")
                st.info(f"🔁 설정 기간 동안 같은 대상의 예상 방문 횟수: {repeat_text}")

        # 결과 완성 안내 — 특정 노선을 대표로 크게 보여주지 않고, 생성된 노선 수만큼
        # 아래 지도 그리드에서 전부 동일하게 보여준다(노선 2 이상이 묻히지 않도록).
        if route_results:
            st.markdown(
                f"""
                <div style="margin:1rem 0 1.1rem;padding:1.25rem 1.4rem;border-radius:18px;
                            background:linear-gradient(135deg,#e9f8ef 0%,#f5fbf7 55%,#fff8df 100%);
                            border:1px solid #9bcdb0;box-shadow:0 10px 28px -18px rgba(25,100,62,.65);">
                  <div style="font-size:1.45rem;font-weight:850;color:#165d39;margin-bottom:.3rem;">
                    🎉 노선이 완성되었습니다!
                  </div>
                  <div style="font-size:1rem;font-weight:650;color:#29483a;">
                    아래에서 생성된 {len(route_results)}개 노선을 지도로 모두 확인하세요.
                    지도 아래 '상세보기'를 열면 방문순서·카카오맵·QR코드가 나옵니다.
                  </div>
                </div>
                """,
                unsafe_allow_html=True,
            )
            api_calls_used = int(meta.get("api_calls_used", 0))
            api_call_limit = int(meta.get("api_call_limit", API_CALL_LIMIT))
            validated_target_count = int(meta.get("validated_target_count", 217))
            st.markdown(
                f"""
                <div style="margin:-.2rem 0 1.25rem;padding:1rem 1.2rem;border-radius:14px;
                            background:#f5f8ff;border:1px solid #b8c8e8;
                            box-shadow:0 8px 20px -18px rgba(31,71,135,.7);">
                  <div style="font-size:.92rem;font-weight:750;color:#284f84;margin-bottom:.35rem;">
                    ⚙️ 대규모 데이터 처리 규모
                  </div>
                  <div style="font-size:1.12rem;font-weight:800;color:#17375f;">
                    이번 작업 사용 API 호출 <span style="color:#126f4b;">{api_calls_used:,}건</span>
                    <span style="color:#68788d;font-weight:650;"> / 일일 기준 참고 한도 {api_call_limit:,}건</span>
                  </div>
                  <div style="margin-top:.3rem;font-size:.88rem;color:#53657b;">
                    조회 대상이 많으면 한 작업에서도 수백~천 건 이상 사용할 수 있어, 하루 사용량 기준으로 표시합니다.
                  </div>
                </div>
                """,
                unsafe_allow_html=True,
            )

        # ---- 담당 조 · 조원 입력(지리조사는 개인정보 보호를 위해 엑셀에서만 작성) ----
        st.markdown("### 📌 다음 작업")
        is_center_route = meta.get("purpose") == "⑤ 지리조사(센터용)"
        if is_center_route:
            st.caption("개인정보 보호를 위해 담당 조·조원은 화면에서 입력하지 않습니다. 결과자료를 내려받은 뒤 엑셀에서 작성하세요.")
            action_right = st.container()
        else:
            st.caption("담당자를 입력하거나 완성된 결과자료를 내려받으세요.")
            action_left, action_right = st.columns(2)
            with action_left:
                with st.expander("👥 담당 조·조원 입력 (선택)", expanded=False):
                    st.caption("필요한 경우에만 입력하세요. 입력 내용은 최종 엑셀 파일에 반영됩니다.")
                    for rr in route_results:
                        st.markdown(f"**노선 {rr['route_no']}**")
                        st.text_input("담당 조 이름", key=f"team_name_{rr['route_no']}",
                                      placeholder="예) 가천1팀1조", label_visibility="collapsed")
                        st.text_input("조원", key=f"team_members_{rr['route_no']}",
                                      placeholder="조원 예) 홍길동, 이순신", label_visibility="collapsed")

        # ---- 엑셀(xlsx) 다운로드: 노선 1개 = 1행, 경유지가 옆으로 펼쳐지는 가로형 ----
        def build_wide_excel(station, route_results, far_points, meta):
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
            from openpyxl.utils import get_column_letter

            wb = Workbook()
            ws = wb.active
            ws.title = "순찰노선"

            prefix = (meta.get("route_prefix") or "").strip() or station["name"]
            base_cols = ["연번", "부서명", "노선이름", "구분", "담당 조", "조원", "출발지"]
            max_stops = max((len(rr["legs"]) for rr in route_results), default=0)

            thin = Side(style="thin", color="9AA59D")
            border = Border(left=thin, right=thin, top=thin, bottom=thin)
            head_fill = PatternFill("solid", fgColor="F4DDD8")
            head_font = Font(bold=True, size=10)
            center = Alignment(horizontal="center", vertical="center", wrap_text=True)

            # 제목 줄
            total_cols = len(base_cols) + max_stops * 3 + 3
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(total_cols, 1))
            title_parts = [value for value in (
                meta.get("purpose", ""), meta.get("vehicle", ""), meta.get("period", "")
            ) if value]
            title_cell = ws.cell(
                row=1, column=1,
                value=f"{meta.get('title', '순찰노선')}   [{ ' · '.join(title_parts) }]",
            )
            title_cell.font = Font(bold=True, size=13)
            title_cell.alignment = Alignment(horizontal="center", vertical="center")

            HEAD1, HEAD2, DATA0 = 2, 3, 4

            for i, name in enumerate(base_cols, start=1):
                ws.merge_cells(start_row=HEAD1, start_column=i, end_row=HEAD2, end_column=i)
                ws.cell(row=HEAD1, column=i, value=name)

            col = len(base_cols) + 1
            for i in range(1, max_stops + 1):
                ws.merge_cells(start_row=HEAD1, start_column=col, end_row=HEAD1, end_column=col + 2)
                ws.cell(row=HEAD1, column=col, value=f"경유{i}")
                ws.cell(row=HEAD2, column=col, value="대상명(주소)")
                ws.cell(row=HEAD2, column=col + 1, value="거리(km)")
                ws.cell(row=HEAD2, column=col + 2, value="누적(km)")
                col += 3

            for label in ("귀소", "노선거리(km)", "순찰 총 소요시간(분)"):
                ws.merge_cells(start_row=HEAD1, start_column=col, end_row=HEAD2, end_column=col)
                ws.cell(row=HEAD1, column=col, value=label)
                col += 1
            last_col = col - 1

            for r_ in (HEAD1, HEAD2):
                for c_ in range(1, last_col + 1):
                    cell = ws.cell(row=r_, column=c_)
                    cell.fill = head_fill
                    cell.font = head_font
                    cell.alignment = center
                    cell.border = border

            r = DATA0
            for rr in route_results:
                no = rr["route_no"]
                c = 1
                ws.cell(row=r, column=c, value=no); c += 1
                ws.cell(row=r, column=c, value=station["name"]); c += 1
                ws.cell(row=r, column=c, value=f"{prefix}노선{no}"); c += 1
                ws.cell(row=r, column=c, value="근거리"); c += 1
                auto_team = f"{rr.get('vehicle_no')}호차" if rr.get("vehicle_no") else ""
                auto_members = ", ".join(rr.get("assigned_members") or [])
                if meta.get("purpose") == "⑤ 지리조사(센터용)":
                    team_value = ""
                    members_value = ""
                else:
                    team_value = st.session_state.get(f"team_name_{no}", "") or auto_team
                    members_value = st.session_state.get(f"team_members_{no}", "") or auto_members
                ws.cell(row=r, column=c, value=team_value); c += 1
                ws.cell(row=r, column=c, value=members_value); c += 1
                ws.cell(row=r, column=c, value=station["name"]); c += 1
                acc_km = 0.0
                for leg in rr["legs"]:
                    acc_km += leg["km"]
                    ws.cell(row=r, column=c, value=stop_label(leg["to"], leg.get("to_address"))); c += 1
                    ws.cell(row=r, column=c, value=round(leg["km"], 1)); c += 1
                    ws.cell(row=r, column=c, value=round(acc_km, 1)); c += 1
                c += (max_stops - len(rr["legs"])) * 3  # 경유지 수가 적은 노선은 빈칸 패딩
                ws.cell(row=r, column=c, value=station["name"]); c += 1
                ws.cell(row=r, column=c, value=round(rr["total_km"], 1)); c += 1
                ws.cell(row=r, column=c, value=round(rr["total_min"])); c += 1
                r += 1

            for idx, p in enumerate(far_points, start=1):
                c = 1
                ws.cell(row=r, column=c, value=f"장거리{idx}"); c += 1
                ws.cell(row=r, column=c, value=station["name"]); c += 1
                ws.cell(row=r, column=c, value="-"); c += 1
                ws.cell(row=r, column=c, value="원거리"); c += 1
                c += 2  # 담당 조 · 조원 빈칸
                ws.cell(row=r, column=c, value=station["name"]); c += 1
                ws.cell(row=r, column=c, value=stop_label(p["name"], p.get("address"))); c += 1
                ws.cell(row=r, column=c, value=p.get("도로거리_km", "")); c += 1
                r += 1

            for row_cells in ws.iter_rows(min_row=DATA0, max_row=r - 1, min_col=1, max_col=last_col):
                for cell in row_cells:
                    cell.border = border
                    cell.alignment = Alignment(vertical="center", wrap_text=True)

            for i in range(1, last_col + 1):
                width = 24 if (i > len(base_cols) and (i - len(base_cols) - 1) % 3 == 0) else 13
                ws.column_dimensions[get_column_letter(i)].width = width
            ws.freeze_panes = ws.cell(row=DATA0, column=len(base_cols) + 1)

            buf = io.BytesIO()
            wb.save(buf)
            return buf.getvalue()

        def build_center_integrated_excel(station, route_results, meta):
            from openpyxl import Workbook
            from openpyxl.drawing.image import Image as XLImage
            from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
            from openpyxl.utils import get_column_letter

            wb = Workbook()
            ws = wb.active
            ws.title = "지리조사 통합표"
            ws.sheet_view.showGridLines = False

            thin = Side(style="thin", color="D9D9D9")
            border = Border(left=thin, right=thin, top=thin, bottom=thin)
            title_fill = PatternFill("solid", fgColor="17375F")
            header_fill = PatternFill("solid", fgColor="D9EAF7")
            summary_fill = PatternFill("solid", fgColor="F7E9EA")
            center = Alignment(horizontal="center", vertical="center", wrap_text=True)
            left = Alignment(horizontal="left", vertical="center", wrap_text=True)
            right = Alignment(horizontal="right", vertical="center", wrap_text=True)

            columns = [
                "팀", "조", "노선", "순번", "대상명", "주소",
                "구간거리(km)", "누적거리(km)", "노선거리(km)", "총소요시간(분)",
                "카카오맵", "QR",
            ]
            last_col = len(columns)

            title = meta.get("title", "센터 지리조사 노선 결과")
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
            title_cell = ws.cell(row=1, column=1, value=f"{title} - 지도·팀별목록 통합 엑셀")
            title_cell.fill = title_fill
            title_cell.font = Font(color="FFFFFF", bold=True, size=14)
            title_cell.alignment = center

            basis = meta.get("hydrant_distribution_basis") or meta.get("basis") or ""
            ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
            ws.cell(
                row=2, column=1,
                value=(
                    f"출발·복귀: {station['name']} / 분배기준: {basis} / "
                    f"{meta.get('team_info', '')}"
                ),
            ).alignment = left

            summary_start = 4
            ws.cell(row=summary_start, column=1, value="팀별 요약표")
            ws.merge_cells(start_row=summary_start, start_column=1, end_row=summary_start, end_column=last_col)
            ws.cell(row=summary_start, column=1).fill = summary_fill
            ws.cell(row=summary_start, column=1).font = Font(bold=True, size=12)
            ws.cell(row=summary_start, column=1).alignment = center

            summary_headers = ["팀", "조/노선 수", "대상 수", "총 거리(km)", "총 소요시간(분)", "비고"]
            for col_no, header in enumerate(summary_headers, start=1):
                cell = ws.cell(row=summary_start + 1, column=col_no, value=header)
                cell.fill = header_fill
                cell.font = Font(bold=True)
                cell.border = border
                cell.alignment = center

            team_summary = {}
            team_route_order = {}
            for rr in route_results:
                team_no = rr.get("vehicle_no") or rr["route_no"]
                team_summary.setdefault(team_no, {"routes": 0, "stops": 0, "km": 0.0, "min": 0.0})
                team_summary[team_no]["routes"] += 1
                team_summary[team_no]["stops"] += len(rr["stops"])
                team_summary[team_no]["km"] += rr["total_km"]
                team_summary[team_no]["min"] += rr["total_min"]
                team_route_order[rr["route_no"]] = team_summary[team_no]["routes"]

            row = summary_start + 2
            for team_no in sorted(team_summary):
                info = team_summary[team_no]
                values = [
                    f"{team_no}팀",
                    info["routes"],
                    info["stops"],
                    round(info["km"], 1),
                    round(info["min"]),
                    "",
                ]
                for col_no, value in enumerate(values, start=1):
                    cell = ws.cell(row=row, column=col_no, value=value)
                    cell.border = border
                    cell.alignment = center if col_no != 6 else left
                row += 1

            list_start = row + 2
            ws.cell(row=list_start, column=1, value="팀별·조별 배정 목록")
            ws.merge_cells(start_row=list_start, start_column=1, end_row=list_start, end_column=last_col)
            ws.cell(row=list_start, column=1).fill = summary_fill
            ws.cell(row=list_start, column=1).font = Font(bold=True, size=12)
            ws.cell(row=list_start, column=1).alignment = center

            header_row = list_start + 1
            for col_no, header in enumerate(columns, start=1):
                cell = ws.cell(row=header_row, column=col_no, value=header)
                cell.fill = header_fill
                cell.font = Font(bold=True)
                cell.border = border
                cell.alignment = center

            row = header_row + 1
            for rr in route_results:
                team_no = rr.get("vehicle_no") or rr["route_no"]
                group_no = team_route_order.get(rr["route_no"], 1)
                route_links = kakao_route_links(station, rr["legs"])
                first_url = route_links[0][0] if route_links else ""
                acc_km = 0.0
                route_start_row = row

                for stop_no, leg in enumerate(rr["legs"], start=1):
                    acc_km += leg["km"]
                    values = [
                        f"{team_no}팀",
                        f"{group_no}조",
                        f"노선 {rr['route_no']}",
                        stop_no,
                        leg["to"],
                        leg.get("to_address", ""),
                        round(leg["km"], 1),
                        round(acc_km, 1),
                        round(rr["total_km"], 1) if stop_no == 1 else "",
                        round(rr["total_min"]) if stop_no == 1 else "",
                        "카카오맵 열기" if stop_no == 1 and first_url else "",
                        "",
                    ]
                    for col_no, value in enumerate(values, start=1):
                        cell = ws.cell(row=row, column=col_no, value=value)
                        cell.border = border
                        cell.alignment = left if col_no in (5, 6) else center
                    if stop_no == 1 and first_url:
                        link_cell = ws.cell(row=row, column=11)
                        link_cell.hyperlink = first_url
                        link_cell.style = "Hyperlink"
                    row += 1

                if first_url:
                    qr_img = XLImage(io.BytesIO(make_qr_png(first_url)))
                    qr_img.width = 92
                    qr_img.height = 92
                    ws.add_image(qr_img, f"L{route_start_row}")
                    ws.row_dimensions[route_start_row].height = 74

            widths = {
                "A": 10, "B": 8, "C": 11, "D": 8, "E": 28, "F": 48,
                "G": 13, "H": 13, "I": 13, "J": 15, "K": 18, "L": 15,
            }
            for column, width in widths.items():
                ws.column_dimensions[column].width = width

            for row_cells in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=last_col):
                for cell in row_cells:
                    if cell.row not in (1,) and cell.value is not None:
                        cell.border = border
                    if cell.column in (7, 8, 9, 10):
                        cell.alignment = right
            ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

            buf = io.BytesIO()
            wb.save(buf)
            return buf.getvalue()

        safe_title = re.sub(r'[\\/:*?"<>|]', "_", meta.get("title", "순찰노선")) or "순찰노선"

        center_print_bytes = build_center_route_print_html(station, route_results, meta)

        # 공문서 작업과 현장 전달에 필요한 4개 자료를 하나의 ZIP으로 묶는다.
        wide_excel_bytes = build_wide_excel(station, route_results, far_points, meta)
        center_integrated_excel_bytes = (
            build_center_integrated_excel(station, route_results, meta) if is_center_route else None
        )
        route_links_excel_bytes = build_route_links_excel(station, route_results)
        qr_zip_bytes = build_qr_zip(station, route_results)
        printable_qr_html_bytes = build_printable_qr_html(station, route_results, meta)

        center_materials = io.BytesIO()
        if is_center_route:
            with zipfile.ZipFile(center_materials, "w", zipfile.ZIP_DEFLATED) as zf:
                zf.writestr(f"1_{safe_title}_노선지도.html", center_print_bytes)
                zf.writestr(f"2_{safe_title}_팀별조별_통합표_QR포함.xlsx", center_integrated_excel_bytes)

        all_materials = io.BytesIO()
        with zipfile.ZipFile(all_materials, "w", zipfile.ZIP_DEFLATED) as zf:
            zf.writestr(f"1_{safe_title}_최종순찰표.xlsx", wide_excel_bytes)
            zf.writestr(f"2_{safe_title}_카카오맵_경로링크.xlsx", route_links_excel_bytes)
            zf.writestr(f"3_{safe_title}_QR코드_전체.zip", qr_zip_bytes)
            zf.writestr(f"4_{safe_title}_QR인쇄문서.html", printable_qr_html_bytes)

        with action_right:
            with st.expander("📦 최종 결과자료 받기", expanded=False):
                st.markdown("### 공문서·현장 전달 자료")
                st.caption("모든 자료는 한 번에 내려받고, 현장에서는 카카오 경로 링크를 바로 여세요.")

                if is_center_route:
                    st.download_button(
                        "지도·팀별목록 한꺼번에 다운로드", data=center_materials.getvalue(),
                        file_name=f"{safe_title}_지도_팀별목록.zip", mime="application/zip",
                        use_container_width=True,
                    )
                else:
                    st.download_button(
                        "📦 모든 자료 한 번에 내려받기", data=all_materials.getvalue(),
                        file_name=f"{safe_title}_모든자료.zip", mime="application/zip",
                        use_container_width=True,
                    )

                st.markdown("### 📱 휴대폰에서 노선 실행하기")
                st.caption(
                    "PC에서 만든 최종 노선 결과를 휴대폰으로 열어, 경유지가 많은 노선도 구간별 카카오맵 버튼으로 이어서 실행합니다. "
                    "파일 업로드나 노선 생성을 다시 하지 않습니다."
                )
                if st.button(
                    "📲 휴대폰에서 노선 실행하기 QR 만들기",
                    type="primary",
                    use_container_width=True,
                    key="create_route_execution_transfer_qr",
                ):
                    now_timestamp = datetime.now().timestamp()
                    try:
                        source_df = st.session_state.get("browser_restored_df")
                        if source_df is None or source_df.empty:
                            source_df = pd.DataFrame([
                                {"대상명": leg.get("to"), "주소": leg.get("to_address", "")}
                                for route in route_results for leg in route.get("legs", [])
                            ])
                        execution_payload = {
                            **route_execution_content(
                                st.session_state.get("browser_source_name") or safe_title,
                                meta.get("title") or patrol_title,
                                station_query,
                                station_result,
                                minimum_transfer_targets(source_df),
                                st.session_state.get("coords_df"),
                                st.session_state.get("coord_api_calls", 0),
                                station,
                                route_results,
                                far_points,
                                meta,
                            ),
                            "saved_at": now_timestamp,
                            "expires_at": now_timestamp + BROWSER_DRAFT_DAYS * 24 * 60 * 60,
                        }
                        transfer_token, transfer_expires_at = create_mobile_transfer(execution_payload)
                        transfer_url = f"{app_public_url()}/?transfer={quote(transfer_token, safe='')}"
                        st.session_state["route_execution_transfer_qr"] = {
                            "url": transfer_url,
                            "png": make_qr_png(transfer_url),
                            "expires_at": transfer_expires_at,
                        }
                    except (OSError, ValueError) as exc:
                        st.session_state.pop("route_execution_transfer_qr", None)
                        st.error(str(exc))

                route_execution_qr = st.session_state.get("route_execution_transfer_qr")
                if route_execution_qr:
                    remaining_seconds = int(
                        float(route_execution_qr["expires_at"]) - datetime.now().timestamp()
                    )
                    if remaining_seconds > 0:
                        st.success("노선 실행용 QR이 준비되었습니다. 현장 휴대폰으로 촬영하세요.")
                        st.image(
                            route_execution_qr["png"],
                            caption="휴대폰에서 노선 결과 열기",
                            width=260,
                        )
                        st.markdown(
                            "1. QR 촬영 · 파세루 열기  \n"
                            "2. 앱 비밀번호 입력  \n"
                            "3. 화면 상단의 현장 노선 이어가기에서 노선 1-1, 노선 1-2 순서대로 카카오맵 열기"
                        )
                        st.warning(
                            "이 QR은 만든 뒤 10분 이내에 한 번만 사용할 수 있습니다. "
                            "가져온 뒤에는 휴대폰 브라우저에 7일간 보관됩니다."
                        )
                    else:
                        st.session_state.pop("route_execution_transfer_qr", None)
                        st.warning("노선 실행용 QR 유효시간 10분이 지났습니다. 새 QR을 만들어주세요.")


                link_box = (st.popover("🔗 카카오 경로 링크 열기", use_container_width=True)
                            if hasattr(st, "popover") else st.expander("🔗 카카오 경로 링크 열기"))
                with link_box:
                    for rr in route_results:
                        route_links = kakao_route_links(station, rr["legs"])
                        for link_no, (kurl, _origin, _destinations) in enumerate(route_links, start=1):
                            suffix = "" if len(route_links) == 1 else f"-{link_no} ({link_no}/{len(route_links)}구간)"
                            st.link_button(f"🚗 노선 {rr['route_no']}{suffix} 열기", kurl,
                                           use_container_width=True)

                if is_center_route:
                    st.caption("문서를 열어 인쇄하면 노선별로 A4 한 장씩 출력됩니다.")
                else:
                    st.caption("순찰표·경로링크·QR코드·QR 인쇄문서가 들어 있습니다.")

        target_min_ref = meta.get("target_min")
        if target_min_ref:
            over = [r for r in route_results if r["total_min"] > target_min_ref]
            if over:
                over_routes = ", ".join(str(r["route_no"]) for r in over)
                st.markdown(
                    f'''<div style="margin:0.75rem 0 1rem;padding:1rem 1.15rem;
                        border:2px solid #e58a14;border-left:7px solid #d97706;border-radius:10px;
                        background:#fff3df;color:#7c3f00;font-weight:650;line-height:1.6;">
                        <strong style="color:#a94f00;font-size:1.02rem;">⚠️ {target_min_ref:g}분 초과 노선 안내</strong><br>
                        목표시간을 넘는 노선이 {len(over)}개 있습니다. (노선 {over_routes})<br>
                        노선당 방문지 수를 줄이거나 목표시간을 늘려 다시 편성해 주세요.
                    </div>''',
                    unsafe_allow_html=True,
                )
            else:
                st.success(f"⏱ 모든 노선이 목표 {target_min_ref}분 이내입니다.")

        st.markdown("### 🗺️ 노선별 지도 · 상세보기")
        st.caption(
            f"생성된 {len(route_results)}개 노선을 모두 지도로 보여드립니다. "
            "각 지도 아래의 '상세보기'를 열면 방문순서·카카오맵 길안내·QR코드가 나옵니다."
        )

        ROUTE_CARDS_PER_ROW = 2
        for row_start in range(0, len(route_results), ROUTE_CARDS_PER_ROW):
            row_routes = route_results[row_start:row_start + ROUTE_CARDS_PER_ROW]
            route_cols = st.columns(len(row_routes))
            for route_col, rr in zip(route_cols, row_routes):
                with route_col:
                    team_name = st.session_state.get(f"team_name_{rr['route_no']}", "")
                    over_mark = " ⏱초과" if target_min_ref and rr["total_min"] > target_min_ref else ""
                    hydrant_label = ""
                    if meta.get("purpose") == "⑤ 지리조사(센터용)" and rr.get("vehicle_no"):
                        members = ", ".join(rr.get("assigned_members") or [])
                        hydrant_label = f" · {rr['vehicle_no']}호차" + (f" · {members}" if members else "")
                    season_runs = (f" · 기간 중 {rr.get('period_runs', 0)}회 예상"
                                   if meta.get("purpose") == "③ 계절순찰" and rr.get("period_runs") else "")
                    inspect_schedule = ""
                    if meta.get("purpose") == "④ 예방검사" and rr.get("inspection_date"):
                        inspect_schedule = (
                            f" · {rr['inspection_date']:%Y-%m-%d} · {rr.get('inspection_team', 1)}팀"
                        )
                    head = (f"노선 {rr['route_no']}" + inspect_schedule + hydrant_label
                            + (f" · {team_name}" if team_name else ""))
                    st.markdown(f"#### 🚒 {head}{over_mark}")
                    st.caption(
                        f"{len(rr['stops'])}개소 · 총 {rr['total_km']:.1f}km · 약 {rr['total_min']:.0f}분"
                        + season_runs
                    )

                    with st.container(border=True):
                        m = folium.Map(location=[station["lat"], station["lng"]], zoom_start=12)
                        if rr["path"]:
                            folium.PolyLine(rr["path"], color="#a33a3f", weight=4, opacity=0.85).add_to(m)

                        # 방문 순서를 지도 위에 숫자로 표시 (기본 핀 대신 번호 원)
                        for i, leg in enumerate(rr["legs"], start=1):
                            folium.Marker(
                                [leg["lat"], leg["lng"]],
                                tooltip=f"{i}. {leg['to']}",
                                icon=folium.DivIcon(
                                    icon_size=(30, 30), icon_anchor=(15, 15),
                                    html=(
                                        '<div style="background:#1f6fb2;color:#ffffff;'
                                        'width:26px;height:26px;border-radius:50%;'
                                        'border:2px solid #ffffff;box-shadow:0 1px 5px rgba(0,0,0,.45);'
                                        'display:flex;align-items:center;justify-content:center;'
                                        'font-family:sans-serif;font-weight:700;font-size:13px;'
                                        f'line-height:1;">{i}</div>'
                                    ),
                                ),
                            ).add_to(m)

                        # 출발·복귀 지점은 눈에 띄게 빨간 '출발' 표식으로
                        folium.Marker(
                            [station["lat"], station["lng"]], tooltip=f"출발·복귀: {station['name']}",
                            icon=folium.DivIcon(
                                icon_size=(56, 26), icon_anchor=(28, 13),
                                html=(
                                    '<div style="background:#a33a3f;color:#ffffff;'
                                    'padding:3px 8px;border-radius:13px;border:2px solid #ffffff;'
                                    'box-shadow:0 1px 5px rgba(0,0,0,.45);text-align:center;'
                                    'font-family:sans-serif;font-weight:700;font-size:12px;'
                                    'line-height:1.2;white-space:nowrap;">🚒 출발</div>'
                                ),
                            ),
                        ).add_to(m)
                        st_folium(m, height=260, use_container_width=True, key=f"map_{rr['route_no']}")

                        with st.expander("🔎 상세보기 · 방문순서 · 카카오맵 · QR코드", expanded=False):
                            rows = []
                            for i, leg in enumerate(rr["legs"], start=1):
                                row = {
                                    "순번": str(i), "지점": stop_label(leg["to"], leg.get("to_address")),
                                    "구간거리(km)": round(leg["km"], 1), "구간시간(분)": round(leg["min"]),
                                }
                                if meta.get("purpose") == "⑤ 지리조사(센터용)":
                                    row["담당"] = leg.get("assigned_to", "")
                                    row["조사시간(분)"] = leg.get("inspection_min", 0)
                                rows.append(row)
                            rows.append({"순번": "", "지점": f"복귀 ({station['name']})",
                                         "구간거리(km)": round(rr["back_km"], 1),
                                         "구간시간(분)": round(rr["back_min"])})
                            st.dataframe(pd.DataFrame(rows), use_container_width=True, hide_index=True)

                            st.markdown("**🟨 카카오맵 — 구간별 노선 실행과 QR코드**")
                            kakao_links = kakao_route_links(station, rr["legs"])
                            st.download_button(
                                f"🖨 노선 {rr['route_no']} QR 인쇄용 문서",
                                data=build_printable_qr_html(station, [rr], meta),
                                file_name=f"{safe_title}_노선_{rr['route_no']}_QR인쇄.html",
                                mime="text/html",
                                key=f"route_print_{rr['route_no']}",
                                use_container_width=True,
                            )
                            for li, (kurl, origin, destinations) in enumerate(kakao_links, start=1):
                                suffix = "" if len(kakao_links) == 1 else f" ({li}/{len(kakao_links)}구간)"
                                seq = " → ".join([origin["name"]] + [p["name"] for p in destinations])
                                qr_png = make_qr_png(kurl)

                                st.link_button(
                                    f"🚗 카카오맵 전체 코스 길안내{suffix}",
                                    kurl,
                                    use_container_width=True,
                                )
                                qr_box = (st.popover(f"📱 QR코드 보기{suffix}", use_container_width=True)
                                          if hasattr(st, "popover")
                                          else st.expander(f"📱 QR코드 보기{suffix}"))
                                with qr_box:
                                    st.image(qr_png, caption="휴대폰 카메라로 스캔하세요.", width=220)
                                    st.warning(
                                        "지도와 예상거리·시간은 입력한 출발지 기준입니다. "
                                        "카카오맵 실행 후 실제 내비는 휴대폰 현재 위치 기준으로 다시 안내될 수 있습니다."
                                    )
                                    st.download_button(
                                        "QR코드 이미지 저장",
                                        data=qr_png,
                                        file_name=f"노선_{rr['route_no']}_카카오맵_QR_{li}.png",
                                        mime="image/png",
                                        key=f"kakao_qr_{rr['route_no']}_{li}",
                                        use_container_width=True,
                                    )
                                st.caption(f"경로: {seq}")

                            if len(kakao_links) > 1:
                                st.caption(
                                    f"※ 카카오맵은 현장 실행 안정성을 위해 한 구간을 최대 {KAKAO_MAX_ROUTE_POINTS}개 지점으로 나눕니다. "
                                    "앞 구간의 마지막 지점을 다음 구간의 출발점으로 겹쳐 이어갑니다. 현장에서 순서대로 열어 주세요."
                                )
                            else:
                                st.caption(
                                    "※ 소방서를 출발지와 최종 목적지로, 순찰 대상을 경유지 순서대로 "
                                    "입력한 카카오맵 링크입니다. QR을 스캔하면 같은 코스가 열립니다."
                                )

                            with st.expander("지점별 개별 길안내(카카오맵) — 예비용", expanded=False):
                                for i, leg in enumerate(rr["legs"], start=1):
                                    st.markdown(
                                        f"- [{i}. {leg['to']} 길안내]({kakao_url(leg['to'], leg['lat'], leg['lng'])})"
                                    )
                                st.markdown(
                                    f"- [🚒 {station['name']} 귀소 길안내]"
                                    f"({kakao_url(station['name'], station['lat'], station['lng'])})"
                                )
                                st.caption(
                                    "전체 코스가 특정 휴대폰에서 열리지 않을 때 사용하세요. 각 버튼은 "
                                    "현재 위치에서 선택한 다음 지점까지 안내합니다."
                                )

        if far_points:
            is_season_far = meta.get("purpose") == "③ 계절순찰"
            st.header("🚙 편도 기준 초과 대상" if is_season_far else "⚠️ 장거리 별도 대상")
            far_df = pd.DataFrame(far_points)
            far_columns = ["name", "address", "도로거리_km"]
            if is_season_far:
                far_columns += ["편도시간_분", "권장수행"]
                st.info("선택한 편도 시간 또는 거리 제한값을 넘는 대상입니다. "
                        "수행자·차량·제한 기준을 조정하거나 별도 순찰 대상으로 검토하세요.")
            st.dataframe(far_df[[column for column in far_columns if column in far_df.columns]],
                         use_container_width=True, hide_index=True)
            for p in far_points:
                st.markdown(
                    f"- [{p['name']} 길안내]({kakao_url(p['name'], p['lat'], p['lng'])})"
                    f" · 실도로거리 {p['도로거리_km']}km"
                    + (f" · 편도 약 {p.get('편도시간_분', '-')}분 · {p.get('권장수행', '')}"
                       if is_season_far else "")
                )

        st.divider()
        st.caption(
            "⚠️ 이 페이지의 API 키는 서버(Secrets)에만 저장되며 브라우저로 노출되지 않습니다. "
            "실제 도로거리·소요시간은 NCP Geocoding·Directions5 실시간 계산 결과입니다. "
            "📲 휴대폰에서는 브라우저 메뉴의 '홈 화면에 추가'를 누르면 앱처럼 사용할 수 있습니다."
        )

    st.info(
        "📌 본 앱이 생성한 순찰노선은 업무 지원을 위한 참고자료입니다. "
        "현장 여건, 도로 상황, 출동 공백 및 대상별 특성을 담당자가 충분히 검토한 후 사용하시기 바랍니다. "
        "최종 노선의 검토·결정과 운용 책임은 프로그램 사용자 및 해당 부서에 있습니다."
    )

    st.markdown(
        '<div style="margin-top:18px;padding:14px 8px 4px;text-align:center;'
        'border-top:1px solid #d9dee4;color:#5f6b64;font-size:0.86rem;">'
        '파세루 오리진&nbsp; | &nbsp;기획·개발 임미성&nbsp; | &nbsp;문의: '
        '<a href="mailto:emtmisung@gmail.com" style="color:#1f6fb2;'
        'text-decoration:none;">emtmisung@gmail.com</a></div>',
        unsafe_allow_html=True,
    )
    return locals()
