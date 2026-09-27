"""대상 목록 파일(CSV·XLSX·XLS·HWPX) 읽기와 빈 양식."""
import io
import re
import zipfile

import pandas as pd

from paseru.settings import SAMPLE_TARGETS


# ----------------------------------------------------------------------------
# 한글(hwpx) 표 파싱 — 데모(웹 프로토타입)와 동일한 방식
# ----------------------------------------------------------------------------
HEADER_WORDS = re.compile(r"^(연번|no\\.?|번호|구분|이름|명칭|대상|대상명|대상물명|시설명|소화전명|소화전번호|관리번호|대상명주소|주소|주소지|소재지|정제_주소|비고)$", re.I)

NAME_HEADER_WORDS = (
    "대상물명", "대상명", "시설명", "소화전명", "소화전번호", "관리번호",
    "시설물명", "장소명", "명칭", "이름", "대상",
)
ADDRESS_HEADER_WORDS = ("주소", "주소지", "소재지", "위치", "설치위치")
SERIAL_HEADER_WORDS = ("연번", "순번", "번호", "no")


def _header_text(value):
    """헤더 비교용 문자열. 공백·줄바꿈·밑줄 차이는 무시한다."""
    if pd.isna(value):
        return ""
    return re.sub(r"[\s_]", "", str(value)).lower()


def _cell_text(value):
    """엑셀 빈칸/nan 값을 화면과 검색에서 실제 값처럼 쓰지 않도록 정리한다."""
    if pd.isna(value):
        return ""
    text = str(value).strip()
    return "" if text.lower() in ("nan", "none", "nat") else text


def _looks_like_address_text(value):
    text = _cell_text(value)
    if not text:
        return False
    lowered = text.lower()
    if any(word in lowered for word in ("조회일시", "검색일시", "시도내역", "실패사유", "비고")):
        return False
    has_place_word = bool(re.search(r"(시|군|구|읍|면|동|리|로|길|번지|산\s*\d|경상|전라|충청|강원|경기|서울|부산|대구|인천|광주|대전|울산|세종|제주)", text))
    has_number = bool(re.search(r"\d", text))
    return has_place_word and has_number


def _address_column_score(series):
    values = list(series.dropna()) if series is not None else []
    nonempty = [_cell_text(v) for v in values if _cell_text(v)]
    if not nonempty:
        return 0
    return sum(1 for value in nonempty if _looks_like_address_text(value))


def _find_header_row(raw_df, scan_rows=80):
    """제목행이 위에 있어도 대상명·주소가 있는 실제 헤더행을 찾는다."""
    for row_idx in range(min(scan_rows, len(raw_df))):
        tokens = [_header_text(v) for v in raw_df.iloc[row_idx].tolist()]
        tokens = [v for v in tokens if v]
        has_name = any(any(word in token for word in NAME_HEADER_WORDS) for token in tokens)
        has_address = any(any(word in token for word in ADDRESS_HEADER_WORDS) for token in tokens)
        has_serial = any(token in SERIAL_HEADER_WORDS for token in tokens)
        if has_address and (has_name or has_serial):
            return row_idx
    return None


def _looks_like_headerless_data(raw_df):
    """항목명까지 지운 파일에서 첫 실제 대상을 헤더로 잃지 않도록 판별한다."""
    if raw_df is None or raw_df.empty:
        return False
    values = [v for v in raw_df.iloc[0].tolist() if not pd.isna(v) and str(v).strip()]
    if len(values) < 2:
        return False
    first = str(values[0]).strip()
    if re.fullmatch(r"\d+(?:\.0)?", first):
        return True
    return any(re.search(r"(?:시|군|구|읍|면|동|리)\s*\S*\d", str(v)) for v in values)


def _make_unique_headers(values):
    headers, seen = [], {}
    for idx, value in enumerate(values, start=1):
        base = str(value).strip() if not pd.isna(value) and str(value).strip() else f"열{idx}"
        seen[base] = seen.get(base, 0) + 1
        headers.append(base if seen[base] == 1 else f"{base}_{seen[base]}")
    return headers


def normalize_uploaded_table(raw_df):
    """제목행/헤더행/헤더 없는 목록을 모두 실제 대상 행 기준으로 정리한다."""
    if raw_df is None or raw_df.empty:
        return raw_df

    raw_df = raw_df.dropna(how="all").dropna(axis=1, how="all").reset_index(drop=True)
    header_row = _find_header_row(raw_df)
    if header_row is not None:
        df = raw_df.iloc[header_row + 1:].copy()
        df.columns = _make_unique_headers(raw_df.iloc[header_row].tolist())
    elif _looks_like_headerless_data(raw_df):
        df = raw_df.copy()
        width = len(df.columns)
        first_value = str(df.iloc[0, 0]).strip() if width else ""
        if width >= 3 and re.fullmatch(r"\d+(?:\.0)?", first_value):
            defaults = ["연번", "대상명", "주소", "비고", "위도", "경도"]
        else:
            defaults = ["대상명", "주소", "비고", "위도", "경도"]
        df.columns = defaults[:width] + [f"열{i}" for i in range(len(defaults) + 1, width + 1)]
    else:
        # 기존 방식과의 호환: 첫 행을 일반적인 열 이름으로 사용한다.
        df = raw_df.iloc[1:].copy()
        df.columns = _make_unique_headers(raw_df.iloc[0].tolist())

    df = df.dropna(axis=1, how="all").dropna(how="all").reset_index(drop=True)
    # 파일 중간에 항목명이 반복된 경우 대상 건수에서 제외한다.
    repeated_headers = df.apply(
        lambda row: _find_header_row(pd.DataFrame([row.tolist()]), scan_rows=1) == 0,
        axis=1,
    )
    df = df.loc[~repeated_headers].reset_index(drop=True)
    if len(df.columns):
        name_index = find_name_column_index(list(df.columns))
        address_index = find_address_column_index(list(df.columns), name_index, df)
        keep_rows = df.apply(
            lambda row: bool(_cell_text(row.iloc[name_index]) or _cell_text(row.iloc[address_index])),
            axis=1,
        )
        df = df.loc[keep_rows].reset_index(drop=True)
    return df


def read_uploaded_table(file_bytes, file_name):
    """CSV/엑셀을 헤더 지정 없이 먼저 읽은 뒤 실제 헤더행을 자동 판별한다."""
    source = io.BytesIO(file_bytes)
    if file_name.lower().endswith(".csv"):
        try:
            raw_df = pd.read_csv(source, header=None)
        except UnicodeDecodeError:
            source.seek(0)
            raw_df = pd.read_csv(source, header=None, encoding="cp949")
    else:
        raw_df = pd.read_excel(source, header=None)
    return normalize_uploaded_table(raw_df)


def load_sample_targets():
    """예시 원본에서 업무 구분과 출발부서를 제외한 평가용 대상 18곳만 만든다."""
    return pd.DataFrame(SAMPLE_TARGETS, columns=["대상명", "주소"])


def find_name_column_index(columns):
    """순번·연번 대신 실제 대상물명 열을 우선 선택한다."""
    normalized = [_header_text(column) for column in columns]
    for preferred in NAME_HEADER_WORDS:
        for idx, token in enumerate(normalized):
            if preferred in token:
                return idx

    excluded_words = ADDRESS_HEADER_WORDS + ("비고", "조별", "위도", "경도", "lat", "lng", "lon")
    for idx, token in enumerate(normalized):
        if token in SERIAL_HEADER_WORDS:
            continue
        if not any(word in token for word in excluded_words):
            return idx
    return 1 if len(columns) > 1 else 0


def find_address_column_index(columns, name_index, df=None):
    """대상명 열과 겹치지 않는 주소 열을 선택한다. 실제 주소값이 많은 열을 우선한다."""
    normalized = [_header_text(column) for column in columns]

    def score(idx):
        if idx == name_index:
            return -1
        token = normalized[idx]
        header_score = 0
        if "정제" in token and any(word in token for word in ADDRESS_HEADER_WORDS):
            header_score += 30
        elif any(word in token for word in ADDRESS_HEADER_WORDS):
            header_score += 20
        if any(word in token for word in ("조회", "검색", "상태", "결과", "비고", "일시")):
            header_score -= 40
        data_score = 0
        if df is not None and idx < len(df.columns):
            data_score = _address_column_score(df.iloc[:, idx]) * 5
        return header_score + data_score

    if not columns:
        return 0
    candidates = [idx for idx in range(len(columns)) if idx != name_index]
    best = max(candidates, key=score) if candidates else 0
    if score(best) > 0:
        return best

    address = next(
        (idx for idx, token in enumerate(normalized)
         if idx != name_index and any(word in token for word in ADDRESS_HEADER_WORDS)),
        None,
    )
    return address if address is not None else min(3, len(columns) - 1)




def _clean_xml_text(s: str) -> str:
    s = re.sub(r"<[^>]+>", "", s)
    return (s.replace("&lt;", "<").replace("&gt;", ">").replace("&amp;", "&")
             .replace("&quot;", '"').replace("&apos;", "'").strip())


def parse_hwpx(file_bytes: bytes):
    """hwpx 안의 표(또는 문단)를 읽어 DataFrame으로 반환."""
    rows = []
    with zipfile.ZipFile(io.BytesIO(file_bytes)) as zf:
        for name in zf.namelist():
            if not re.search(r"Contents/section\d*\.xml$", name, re.I):
                continue
            xml = zf.read(name).decode("utf-8", errors="ignore")
            table_rows = re.findall(r"<hp:tr[\s>][\s\S]*?</hp:tr>", xml)
            if table_rows:
                for row_xml in table_rows:
                    cells = re.findall(r"<hp:tc[\s>][\s\S]*?</hp:tc>", row_xml)
                    cols = [_clean_xml_text("".join(re.findall(r"<hp:t[^>]*>([\s\S]*?)</hp:t>", c)))
                            for c in cells]
                    if any(cols):
                        rows.append(cols)
            else:
                for chunk in xml.split("<hp:p")[1:]:
                    text = _clean_xml_text("".join(re.findall(r"<hp:t[^>]*>([\s\S]*?)</hp:t>", chunk)))
                    if text:
                        rows.append([text])

    rows = [r for r in rows if r and not HEADER_WORDS.match((r[0] or "").strip())]
    if not rows:
        return None

    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    default_names = ["연번", "주소지", "비고", "정제_주소", "위도(Latitude)", "경도(Longitude)"]
    cols = default_names[:width] + [f"열{i}" for i in range(len(default_names) + 1, width + 1)]
    return pd.DataFrame(rows, columns=cols[:width])


def build_upload_template():
    """대상 목록을 일정한 열 이름으로 작성할 수 있는 빈 엑셀 양식."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "대상목록"
    ws.sheet_view.showGridLines = False

    headers = ["연번", "대상명", "주소", "비고", "위도(선택)", "경도(선택)"]
    ws.append(headers)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    required_fill = PatternFill("solid", fgColor="FFF2CC")
    thin = Side(style="thin", color="D9D9D9")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # 첫 입력행은 비워두되 필수 입력칸을 연한 노랑으로 표시한다.
    for row_no in range(2, 102):
        for col_no in range(1, 7):
            cell = ws.cell(row=row_no, column=col_no)
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.cell(row=row_no, column=2).fill = required_fill
        ws.cell(row=row_no, column=3).fill = required_fill

    widths = {"A": 9, "B": 28, "C": 52, "D": 28, "E": 16, "F": 16}
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    ws.row_dimensions[1].height = 27
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:F101"

    guide = wb.create_sheet("작성안내")
    guide.sheet_view.showGridLines = False
    guide.append(["항목", "필수 여부", "작성 방법"])
    guide_rows = [
        ["대상명", "필수", "시설명 또는 점검 대상명을 입력합니다."],
        ["주소", "필수", "지오코딩할 도로명주소 또는 지번주소를 입력합니다."],
        ["연번", "선택", "자동으로 표시됩니다. 직접 수정해도 됩니다."],
        ["비고", "선택", "노선 편성에 필요한 일반 참고사항만 입력합니다."],
        ["위도·경도", "선택", "이미 검증한 좌표가 있을 때만 입력합니다. 없으면 비워두세요."],
        ["개인정보", "입력 금지", "성명, 전화번호, 주민등록번호, 검사결과 등은 입력하지 않습니다."],
    ]
    for row in guide_rows:
        guide.append(row)
    for cell in guide[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in guide.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = Border(bottom=thin)
    guide.column_dimensions["A"].width = 18
    guide.column_dimensions["B"].width = 14
    guide.column_dimensions["C"].width = 72
    guide.freeze_panes = "A2"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
